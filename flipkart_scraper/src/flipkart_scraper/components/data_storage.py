import pandas as pd
import os
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

class ExcelStorage:
    def __init__(self, filename='Flipkart_WashingMachines.xlsx'):
        self.filename = os.path.join('data', filename)
        os.makedirs(os.path.dirname(self.filename), exist_ok=True)

    def store(self, df, brand):
        """
        Stores the DataFrame in an Excel file. Appends to existing sheets without overwriting.
        """
        if df.empty:
            return

        if not os.path.exists(self.filename):
            with pd.ExcelWriter(self.filename, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name=brand, index=False)
        else:
            book = load_workbook(self.filename)
            if brand in book.sheetnames:
                sheet = book[brand]
                # Append rows without the header
                for row in dataframe_to_rows(df, index=False, header=False):
                    sheet.append(row)
                book.save(self.filename)
            else:
                # If sheet does not exist, create it with header
                with pd.ExcelWriter(self.filename, engine='openpyxl', mode='a') as writer:
                    df.to_excel(writer, sheet_name=brand, index=False)

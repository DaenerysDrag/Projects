import pandas as pd
import os
import logging
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows


class ExcelStorage:
    def __init__(self, filename='Flipkart_WashingMachines.xlsx'):
        self.filename = os.path.join('data', filename)
        os.makedirs(os.path.dirname(self.filename), exist_ok=True)

    def store(self, df, brand):
        """
        Append new rows for *brand* into the Excel file.
        - Creates the file / sheet if they don't exist yet.
        - Skips rows whose URL already exists in the sheet (deduplication).
        """
        if df.empty:
            logging.info(f"No data to store for brand '{brand}'.")
            return

        # ---- Brand-new file ------------------------------------------------
        if not os.path.exists(self.filename):
            with pd.ExcelWriter(self.filename, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name=brand, index=False)
            logging.info(
                f"Created '{self.filename}' with sheet '{brand}' "
                f"({len(df)} row(s))."
            )
            return

        # ---- File exists ----------------------------------------------------
        book = load_workbook(self.filename)

        if brand in book.sheetnames:
            sheet = book[brand]

            # Collect existing URLs from the sheet to detect duplicates
            header = [cell.value for cell in sheet[1]]
            url_col_idx = header.index('URL') if 'URL' in header else None

            existing_urls: set = set()
            if url_col_idx is not None:
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    val = row[url_col_idx]
                    if val:
                        existing_urls.add(val)

            # Keep only rows that are genuinely new
            new_df = df[~df['URL'].isin(existing_urls)] if existing_urls else df

            if new_df.empty:
                logging.info(
                    f"All {len(df)} scraped row(s) for '{brand}' already exist — nothing appended."
                )
                return

            for row in dataframe_to_rows(new_df, index=False, header=False):
                sheet.append(row)
            book.save(self.filename)
            logging.info(
                f"Appended {len(new_df)} new row(s) for '{brand}' "
                f"({len(df) - len(new_df)} duplicate(s) skipped)."
            )

        else:
            # Sheet for this brand doesn't exist yet — create it
            with pd.ExcelWriter(self.filename, engine='openpyxl', mode='a') as writer:
                df.to_excel(writer, sheet_name=brand, index=False)
            logging.info(
                f"Added new sheet '{brand}' to '{self.filename}' ({len(df)} row(s))."
            )
